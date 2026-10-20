#!/usr/bin/env python3
"""Editable Word printables for the sharing unit, generated from data/units.json.
   python3 tools/build_printables.py        -> printables/*.docx  (linked from teacher.html)
Sheets are real tables: change the words, resize a box, recolour a cell in Word as usual."""
import random, pathlib, sys
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from sharing_lib import *

OUT = ROOT / 'printables'; OUT.mkdir(exist_ok=True)
PER_PAGE = 5
BOX = 2.3          # cm; one box per sound
GAP = 0.25         # cm spacer row between words
ROW = BOX

def pages(ids, n=PER_PAGE):
    # balanced pages: 10 -> 5+5, 9 -> 5+4, 7 -> 4+3
    k = -(-len(ids) // n); size = -(-len(ids) // k)
    return [ids[i:i + size] for i in range(0, len(ids), size)]

GAPW = 0.7         # cm; the gap between syllables / words (can | dy, ice | cream)

def sheet(u, mode):
    """One small table per word, so each row has its own column widths: boxes, then a narrow gap at each
    syllable or word break, as on the hand-made sheets."""
    doc = landscape_doc(); ids = u['words']
    left = {'A': [4.6, 2.9, 2.2], 'B': [3.1, 6.4], 'C': [1.8]}[mode]
    def fit(w):
        sl = slots(w); nb = sum(1 for g in sl if g); ng = len(sl) - nb
        return (26.7 - sum(left) - ng * GAPW) / nb
    box = round(min([BOX] + [fit(w) for w in ids]), 2)
    n = 0
    for pi, chunk in enumerate(pages(ids)):
        t = title_par(doc, u['title'], u['zh'])
        if pi: t.paragraph_format.page_break_before = True
        for ri, wid in enumerate(chunk):
            if ri: tiny_par(doc, height_pt=GAP * 28.35)
            sl = slots(wid); n += 1; w = D['words'][wid]
            tbl = doc.add_table(rows=1, cols=len(left) + len(sl)); fixed_table(tbl)
            set_widths(tbl, left + [box if g else GAPW for g in sl])
            row = tbl.rows[0]; cs = row.cells
            for c in cs: no_borders(c); margins(c, 0.05)
            row_height(row, box)
            if mode == 'A':
                for c in cs[:3]: shade(c, GREEN)
                write(cs[0], wid, 26, bold=True); write(cs[1], w['zh'], 24, zh=True); add_picture(cs[2], w['icon'], 1.3, 1.8)
                borders(cs[0], right=('dashed', 8, '808080'))   # fold line: the English folds under
            elif mode == 'B':
                for c in cs[:2]: shade(c, GREEN)
                write(cs[0], w['zh'], 26, zh=True); write(cs[1], w['def'], 13, align='left')
            else:
                shade(cs[0], GREEN); write(cs[0], str(n), 24, bold=True)
            for k, g in enumerate(sl):
                if g is None: continue
                c = cs[len(left) + k]; shade(c, GREY); box_borders(c)
        # the sheet is its own explanation: no instruction text
    return doc

def tiny_par(doc, break_before=False, height_pt=1):
    """A small paragraph: Word needs one between tables, and it must not push a page over."""
    from docx.enum.text import WD_LINE_SPACING
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0); pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY; pf.line_spacing = Pt(height_pt)
    pf.page_break_before = break_before
    return p

def cards(u):
    doc = landscape_doc(); ids = u['words']; rnd = random.Random(7)
    kinds = [('word', ids), ('zh', rnd.sample(ids, len(ids))), ('pic', rnd.sample(ids, len(ids))), ('def', rnd.sample(ids, len(ids)))]
    for pi, (kind, order) in enumerate(kinds):
        tiny_par(doc, break_before=pi > 0)
        tbl = doc.add_table(rows=3, cols=4); fixed_table(tbl); set_widths(tbl, [6.675] * 4)
        for i, c in enumerate([c for r in tbl.rows for c in r.cells]):
            borders(c, top=('dashed', 6, '808080'), bottom=('dashed', 6, '808080'), left=('dashed', 6, '808080'), right=('dashed', 6, '808080')); margins(c, 0.15)
            if i < len(order):
                w = D['words'][order[i]]
                if kind == 'word': write(c, order[i], 28, bold=True)
                elif kind == 'zh': write(c, w['zh'], 34, zh=True)
                elif kind == 'pic': add_picture(c, w['icon'], 2.6)
                else: write(c, w['def'], 15)
        for r in tbl.rows: row_height(r, 5.9)
    return doc

def rhyme(u):
    doc = landscape_doc(); fams = [f for f in D['rhymes'] if f['unit'] == u['id']]
    for pi, f in enumerate(fams):
        onset, rime = key_parts(f); w = D['words'][f['key']]; nb = len(onset) + len(rime)
        t = title_par(doc, '-' + f['rime'], None, 22)
        if pi: t.paragraph_format.page_break_before = True
        maxs = 5; left = [4.6, 2.9, 2.2]
        tbl = doc.add_table(rows=9, cols=3 + maxs); fixed_table(tbl); set_widths(tbl, left + [BOX] * maxs)
        for ri in range(9):
            row = tbl.rows[ri]
            for c in row.cells: no_borders(c); margins(c, 0.05)
            if ri % 2: row_height(row, GAP); continue
            row_height(row, BOX); cs = row.cells
            if ri == 0:
                for c in cs[:3]: shade(c, GREEN)
                write(cs[0], f['key'], 26, bold=True); write(cs[1], w['zh'], 24, zh=True); add_picture(cs[2], w['icon'], 1.3, 1.8)
                for k, g in enumerate(onset + rime):
                    c = cs[3 + k]; shade(c, GREY); box_borders(c); write(c, g, 28, bold=True)
            else:
                m = cs[0].merge(cs[2]); shade(m, GREEN)
                for k in range(nb):
                    c = cs[3 + k]; box_borders(c)
                    if k < len(onset): shade(c, 'FFFFFF')            # write the new start here
                    else: shade(c, GREY); write(c, rime[k - len(onset)], 28, bold=True)
    return doc

if __name__ == '__main__':
    n = 0
    for i, u in enumerate(D['units'], 1):
        base = f'U{i}'
        for mode, name in (('A', '1-boxes-word-and-chinese'), ('B', '2-boxes-meaning-and-chinese'), ('C', '3-boxes-only')):
            sheet(u, mode).save(OUT / f'{base}_{name}.docx'); n += 1
        cards(u).save(OUT / f'{base}_match-cards.docx'); n += 1
        rhyme(u).save(OUT / f'{base}_rhyme-sheets.docx'); n += 1
    print('wrote', n, 'files to', OUT)
