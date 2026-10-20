"""Shared helpers for the printable generators (build_printables.py, build_slides.py): data, emoji icons, box slots, python-docx cell helpers."""
import json, pathlib
from PIL import Image, ImageDraw, ImageFont
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = pathlib.Path(__file__).resolve().parent.parent
D = json.loads((ROOT / 'data' / 'units.json').read_text(encoding='utf-8'))
EMOJI_FONT = '/usr/share/fonts/truetype/noto/NotoColorEmoji.ttf'
CACHE = pathlib.Path(__file__).resolve().parent / '_icons'
EN_FONT, ZH_FONT = 'Arial', 'Microsoft JhengHei'
GREEN, GREY = 'E9EFDC', 'D9D9D9'     # the pale green / grey from the hand-made sheets

def icon_png(icon):
    """Render an emoji (or a short run of them) to a transparent PNG; cached."""
    CACHE.mkdir(exist_ok=True)
    out = CACHE / (''.join(f'{ord(c):x}-' for c in icon) + '.png')
    if out.exists():
        return str(out)
    f = ImageFont.truetype(EMOJI_FONT, 109)
    chars = list(icon)
    im = Image.new('RGBA', (140 * len(chars), 140), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    for i, c in enumerate(chars):
        d.text((140 * i, 4), c, font=f, embedded_color=True)
    im.crop(im.getbbox()).save(out)
    return str(out)

def slots(word_id):
    """['box' chunk | None for a gap] in order; a gap sits between the words of a phrase."""
    out = []
    for i, part in enumerate(D['words'][word_id]['parts']):
        if i: out.append(None)
        out.extend(part)
    return out

def key_parts(fam):
    """(onset chunks, rime chunks) of a rhyme family's key word, using the word's own sound boxes."""
    parts = D['words'][fam['key']]['parts'][0]
    n = len(fam['key']) - len(fam['rime']); used = i = 0
    while used < n and i < len(parts):
        used += len(parts[i]); i += 1
    return parts[:i], parts[i:]

def unit(uid):
    return next(u for u in D['units'] if u['id'] == uid)

# ---- python-docx cell helpers ----
def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')): tcPr.remove(old)
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexfill); tcPr.append(shd)

def borders(cell, **sides):
    """borders(cell, top=(style,size,color), ...) ; style 'single'/'dashed'/'nil'"""
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None: b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for side, spec in sides.items():
        e = b.find(qn('w:' + side))
        if e is None: e = OxmlElement('w:' + side); b.append(e)
        style, sz, color = spec
        e.set(qn('w:val'), style); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)

def box_borders(cell, sz=18):
    borders(cell, top=('single', sz, '000000'), bottom=('single', sz, '000000'), left=('single', sz, '000000'), right=('single', sz, '000000'))

def no_borders(cell):
    borders(cell, top=('nil', 0, 'auto'), bottom=('nil', 0, 'auto'), left=('nil', 0, 'auto'), right=('nil', 0, 'auto'))

def margins(cell, cm=0.1):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement('w:tcMar')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + side); e.set(qn('w:w'), str(int(cm * 567))); e.set(qn('w:type'), 'dxa'); m.append(e)
    tcPr.append(m)

def row_height(row, cm, exact=True):
    trPr = row._tr.get_or_add_trPr(); h = OxmlElement('w:trHeight')
    h.set(qn('w:val'), str(int(cm * 567))); h.set(qn('w:hRule'), 'exact' if exact else 'atLeast'); trPr.append(h)
    cant = OxmlElement('w:cantSplit'); trPr.append(cant)

def fixed_table(table):
    table.autofit = False
    tblPr = table._tbl.tblPr; lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    # no table borders: every border is set per cell
    tb = OxmlElement('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + side); e.set(qn('w:w'), '0'); e.set(qn('w:type'), 'dxa'); tb.append(e)
    tblPr.append(tb)

def set_widths(table, widths_cm):
    for row in table.rows:
        for c, w in zip(row.cells, widths_cm): c.width = Cm(w)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths_cm): gc.set(qn('w:w'), str(int(w * 567)))

def write(cell, text, size=14, bold=False, zh=False, color=None, align='center', italic=False):
    """Replace a cell's first paragraph with one run; sets Latin and East-Asian fonts."""
    p = cell.paragraphs[0]; p.text = ''
    p.alignment = {'center': WD_ALIGN_PARAGRAPH.CENTER, 'left': WD_ALIGN_PARAGRAPH.LEFT}[align]
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    r.font.name = EN_FONT; rPr = r._r.get_or_add_rPr(); rf = rPr.find(qn('w:rFonts'))
    rf.set(qn('w:eastAsia'), ZH_FONT); rf.set(qn('w:hint'), 'eastAsia')
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return p

def add_picture(cell, icon, height_cm, max_w_cm=None):
    p = cell.paragraphs[0]; p.text = ''; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    path = icon_png(icon); r = p.add_run()
    if max_w_cm:
        w_px, h_px = Image.open(path).size
        if height_cm * w_px / h_px > max_w_cm: r.add_picture(path, width=Cm(max_w_cm)); cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER; return
    r.add_picture(path, height=Cm(height_cm))
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def landscape_doc(margin_cm=1.5):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    doc = Document(); s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = Cm(margin_cm); s.top_margin = s.bottom_margin = Cm(1.2)
    cp = doc.core_properties; cp.author = cp.last_modified_by = ''; cp.title = 'Sharing words'; cp.comments = ''
    st = doc.styles['Normal']; st.font.name = EN_FONT; st.element.rPr.rFonts.set(qn('w:eastAsia'), ZH_FONT); st.font.size = Pt(12)
    return doc

def title_par(doc, text, zh_text=None, size=20):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6); p.paragraph_format.space_before = Pt(0)
    for t, z in ((text, False), (('   ' + zh_text) if zh_text else '', True)):
        if not t: continue
        r = p.add_run(t); r.font.size = Pt(size); r.font.name = EN_FONT
        r._r.get_or_add_rPr().find(qn('w:rFonts')).set(qn('w:eastAsia'), ZH_FONT)
    return p

def page_break(doc):
    from docx.enum.text import WD_BREAK
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE) if doc.paragraphs else None
